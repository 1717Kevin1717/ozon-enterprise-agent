import csv
import hashlib
import io
import math
import re
from pathlib import Path
from typing import Any

from sqlalchemy import desc, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import ROOT
from app.db.models import KnowledgeChunk, KnowledgeDocument

ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".pdf", ".txt", ".md"}
MAX_KNOWLEDGE_BYTES = 10 * 1024 * 1024
EMBEDDING_MODE = "local_hash_embedding_v1"


class KnowledgeError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise KnowledgeError("TEXT_ENCODING_UNSUPPORTED", "文件编码无法识别，请保存为 UTF-8 后重试。")


def _parse_csv(content: bytes) -> str:
    reader = csv.reader(io.StringIO(_decode_text(content)))
    return "\n".join(" | ".join(str(cell).strip() for cell in row) for row in reader if any(str(cell).strip() for cell in row))


def _parse_xlsx(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise KnowledgeError("XLSX_PARSER_UNAVAILABLE", "Excel 解析组件未安装，请安装 requirements.txt 后重试。") from exc
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"工作表：{sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _parse_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise KnowledgeError("PDF_PARSER_UNAVAILABLE", "PDF 解析组件未安装，请安装 requirements.txt 后重试。") from exc
    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
    return "\n".join(pages)


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise KnowledgeError("FILE_TYPE_UNSUPPORTED", "只支持 CSV、Excel、PDF、TXT 和 Markdown 文件。")
    if not content:
        raise KnowledgeError("FILE_EMPTY", "上传文件为空。")
    if len(content) > MAX_KNOWLEDGE_BYTES:
        raise KnowledgeError("FILE_TOO_LARGE", "文件超过 10 MB 限制。")
    text = _parse_csv(content) if extension == ".csv" else _parse_xlsx(content) if extension == ".xlsx" else _parse_pdf(content) if extension == ".pdf" else _decode_text(content)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise KnowledgeError("NO_EXTRACTABLE_TEXT", "文件中没有可提取文本；扫描版 PDF 需要后续 OCR Connector。")
    return text


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 120) -> list[str]:
    paragraphs = [item.strip() for item in re.split(r"\n+", text) if item.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        candidate = f"{current}\n{paragraph}".strip()
        if current and len(candidate) > chunk_size:
            chunks.append(current)
            current = f"{current[-overlap:]}\n{paragraph}".strip()
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def _tokens(text: str) -> list[str]:
    normalized = text.casefold()
    words = re.findall(r"[a-z0-9_]+|[\u4e00-\u9fff]", normalized)
    chinese = "".join(token for token in words if len(token) == 1 and "\u4e00" <= token <= "\u9fff")
    return [token for token in words if len(token) > 1] + [chinese[index:index + 2] for index in range(max(0, len(chinese) - 1))]


def local_embedding(text: str, dimensions: int = 128) -> list[float]:
    vector = [0.0] * dimensions
    for token in _tokens(text):
        digest = hashlib.sha256(token.encode("utf-8")).digest()
        index = int.from_bytes(digest[:4], "big") % dimensions
        vector[index] += -1.0 if digest[4] & 1 else 1.0
    norm = math.sqrt(sum(value * value for value in vector)) or 1.0
    return [round(value / norm, 6) for value in vector]


def cosine(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    return sum(a * b for a, b in zip(left, right))


async def ingest_document(session: AsyncSession, company_id: str, filename: str, content_type: str, content: bytes) -> tuple[KnowledgeDocument, bool]:
    safe_name = Path(filename).name
    checksum = hashlib.sha256(content).hexdigest()
    existing = await session.scalar(select(KnowledgeDocument).where(KnowledgeDocument.company_id == company_id, KnowledgeDocument.checksum == checksum))
    if existing:
        return existing, False
    text = extract_text(safe_name, content)
    chunks = chunk_text(text)
    storage_root = ROOT / "data" / "knowledge" / company_id
    storage_root.mkdir(parents=True, exist_ok=True)
    stored_path = storage_root / f"{checksum[:16]}-{safe_name}"
    stored_path.write_bytes(content)
    document = KnowledgeDocument(company_id=company_id, filename=safe_name, content_type=content_type or "application/octet-stream", storage_path=str(stored_path), checksum=checksum, status="ready", chunk_count=len(chunks), metadata_json={"embedding_mode": EMBEDDING_MODE, "bytes": len(content)})
    session.add(document)
    await session.flush()
    for index, chunk in enumerate(chunks):
        session.add(KnowledgeChunk(company_id=company_id, document_id=document.id, chunk_index=index, content=chunk, embedding=local_embedding(chunk), metadata_json={"embedding_mode": EMBEDDING_MODE, "filename": safe_name}))
    await session.commit()
    await session.refresh(document)
    return document, True


async def list_documents(session: AsyncSession, company_id: str) -> list[KnowledgeDocument]:
    return list((await session.scalars(select(KnowledgeDocument).where(KnowledgeDocument.company_id == company_id).order_by(desc(KnowledgeDocument.created_at)))).all())


async def search_documents(session: AsyncSession, company_id: str, query: str, limit: int) -> list[dict[str, Any]]:
    query_vector = local_embedding(query)
    chunks = list((await session.scalars(select(KnowledgeChunk).where(KnowledgeChunk.company_id == company_id).limit(2000))).all())
    document_ids = {chunk.document_id for chunk in chunks}
    documents = {doc.id: doc for doc in (await session.scalars(select(KnowledgeDocument).where(KnowledgeDocument.company_id == company_id, KnowledgeDocument.id.in_(document_ids)))).all()} if document_ids else {}
    query_terms = set(_tokens(query))
    ranked = []
    for chunk in chunks:
        lexical = len(query_terms & set(_tokens(chunk.content))) / max(len(query_terms), 1)
        vector_score = max(0.0, cosine(query_vector, chunk.embedding or []))
        score = lexical * 0.65 + vector_score * 0.35
        if score > 0:
            document = documents.get(chunk.document_id)
            ranked.append({"chunk_id": chunk.id, "document_id": chunk.document_id, "filename": document.filename if document else "", "content": chunk.content, "score": round(score, 4), "source": "company_knowledge_base", "embedding_mode": EMBEDDING_MODE})
    ranked.sort(key=lambda item: item["score"], reverse=True)
    return ranked[:limit]


def document_view(document: KnowledgeDocument) -> dict[str, Any]:
    return {"id": document.id, "filename": document.filename, "content_type": document.content_type, "source_type": document.source_type, "status": document.status, "chunk_count": document.chunk_count, "error_code": document.error_code, "metadata": document.metadata_json, "created_at": document.created_at, "updated_at": document.updated_at}
